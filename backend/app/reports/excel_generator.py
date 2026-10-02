import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_report(report_data: dict, school_info: dict) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{report_data['month_name']} Attendance"

    # Header styling
    school_name = school_info.get("school_name", "School Attendance System")
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = school_name.upper()
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="1E1B4B")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:J2")
    sub_cell = ws["A2"]
    sub_cell.value = f"MONTHLY ATTENDANCE REPORT — {report_data['month_name']} {report_data['year']} | Class: {report_data['class_name']} - Section {report_data['section_name']}"
    sub_cell.font = Font(name="Calibri", size=11, bold=True, color="475569")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Table Column Headers
    headers = [
        "Roll No", "Student ID", "Admission No", "Student Name",
        "Class", "Section", "Working Days", "Present", "Absent", "Attendance %"
    ]
    ws.append([])  # blank row 3
    ws.append(headers)  # row 4

    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx, col_name in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data Rows
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    low_att_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    for row_idx, s in enumerate(report_data["students"], start=5):
        row_data = [
            s["roll_number"] or "—",
            s["student_code"],
            s["admission_number"],
            s["student_name"],
            report_data["class_name"],
            report_data["section_name"],
            s["working_days"],
            s["present"],
            s["absent"],
            f"{s['percentage']}%",
        ]
        ws.append(row_data)

        use_alt = (row_idx % 2 == 0)
        is_low = s["is_low_attendance"]

        for col_idx in range(1, 11):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_idx != 4 else "left", vertical="center")
            if is_low and col_idx == 10:
                cell.fill = low_att_fill
                cell.font = Font(bold=True, color="DC2626")
            elif use_alt:
                cell.fill = alt_fill

    # Summary Row
    last_row = 5 + len(report_data["students"])
    ws.cell(row=last_row + 1, column=1, value="Class Average Attendance:").font = Font(bold=True)
    ws.cell(row=last_row + 1, column=4, value=f"{report_data['average_attendance']}%").font = Font(bold=True, color="4F46E5")
    ws.cell(row=last_row + 2, column=1, value="Low Attendance Count:").font = Font(bold=True)
    ws.cell(row=last_row + 2, column=4, value=f"{report_data['low_attendance_count']} students").font = Font(bold=True, color="DC2626")

    # Column Auto-width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
